from io import BytesIO

import fitz
from docx import Document as DocxDocument
from fastapi.testclient import TestClient

from backend.app.main import app

client = TestClient(app)
client.post('/api/v1/auth/demo')


def test_health_and_demo_documents():
    assert client.get('/api/v1/health').json()['mode'] == 'demo'
    response = client.get('/api/v1/documents')
    assert response.status_code == 200
    document_ids = {document['id'] for document in response.json()}
    assert {'demo-rental-v1', 'demo-rental-v2'} <= document_ids


def test_grounded_answer_has_verified_citations():
    response = client.post('/api/v1/documents/demo-rental-v1/questions', json={'question': 'When is rent due?'})
    body = response.json()
    assert body['supported'] is True
    assert body['citations'][0]['clause_id'] == 'demo-rental-v1-c2'
    assert body['citations'][0]['document_id'] == 'demo-rental-v1'


def test_unsupported_question_is_honest():
    response = client.post('/api/v1/documents/demo-rental-v1/questions', json={'question': 'What does the statute say?'})
    assert response.json()['supported'] is False
    assert response.json()['citations'] == []


def test_comparison_identifies_real_changes():
    response = client.post('/api/v1/comparisons', json={'left_document_id': 'demo-rental-v1', 'right_document_id': 'demo-rental-v2'})
    assert response.status_code == 200
    assert len(response.json()['changes']) == 2


def test_demo_documents_cannot_be_deleted():
    response = client.delete('/api/v1/documents/demo-rental-v1')
    assert response.status_code == 409


def test_document_operations_require_a_session():
    anonymous = TestClient(app)
    response = anonymous.get('/api/v1/documents')
    assert response.status_code == 401


def test_pdf_upload_extracts_page_aware_clauses():
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), '1. Payment\nMonthly rent is $900 due on the first day.')
    payload = pdf.tobytes()
    response = client.post('/api/v1/documents', files={'file': ('agreement.pdf', payload, 'application/pdf')})
    assert response.status_code == 201
    body = response.json()
    assert body['status'] == 'Extracted'
    assert body['clauses'][0]['page'] == 1
    assert 'Monthly rent' in body['clauses'][1]['text']
    persisted = client.get(f"/api/v1/documents/{body['id']}")
    assert persisted.status_code == 200
    assert persisted.json()['clauses'][1]['text'] == body['clauses'][1]['text']


def test_docx_upload_extracts_paragraphs():
    source = DocxDocument()
    source.add_paragraph('1. Parties')
    source.add_paragraph('The agreement is between Alex and Sam.')
    output = BytesIO()
    source.save(output)
    response = client.post('/api/v1/documents', files={'file': ('agreement.docx', output.getvalue(), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
    assert response.status_code == 201
    assert response.json()['type'] == 'DOCX'
    assert len(response.json()['clauses']) == 2


def test_upload_rejects_invalid_signature():
    response = client.post('/api/v1/documents', files={'file': ('agreement.pdf', b'not a pdf', 'application/pdf')})
    assert response.status_code == 415


def test_upload_rejects_oversized_file():
    response = client.post('/api/v1/documents', files={'file': ('large.pdf', b'x' * (10 * 1024 * 1024 + 1), 'application/pdf')})
    assert response.status_code == 413
